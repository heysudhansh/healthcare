import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_architecture_doc(output_path):
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Arial"
            r_b.font.size = Pt(9.5)
            r_b.font.bold = True
            r_b.font.color.rgb = RGBColor(0, 0, 0)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_bullet(bold_label, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r_b = p.add_run(bold_label)
        r_b.font.name = "Arial"
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0, 0, 0)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_table(headers, data, col_widths):
        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        for c_idx, cell in enumerate(table.rows[0].cells):
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(headers[c_idx])
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
        for r_idx, row_data in enumerate(data):
            row = table.rows[r_idx + 1]
            for c_idx, cell in enumerate(row.cells):
                cell.width = col_widths[c_idx]
                set_cell_margins(cell, top=45, bottom=45, left=70, right=70)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                r = p.add_run(str(row_data[c_idx]))
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0, 0, 0)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("HEALTHCARE SYSTEM - ARCHITECTURE & CODE FLOW GUIDE")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    add_body("Comprehensive technical breakdown of System Architecture, DTO Validation, Global Exception Handling, Service Rules, Frontend Guards, and Unit Testing Strategy.")

    # 1. Architecture
    add_h1("1. SYSTEM OVERVIEW & LAYERED ARCHITECTURE")
    add_body("The application is structured into a classic N-Tier Layered Architecture with clear separation of concerns:")
    add_bullet("1. Frontend Layer (React.js - Port 3000): ", "Handles user interactions, real-time input filtering, and client-side alerts.")
    add_bullet("2. Controller Layer (Spring MVC - Port 8081): ", "Exposes REST endpoints and activates DTO validation via @Valid.")
    add_bullet("3. DTO Layer (Data Transfer Objects): ", "Enforces strict input formatting, syntax constraints, and field-level rules.")
    add_bullet("4. Service Layer (Spring @Service): ", "Enforces core hospital business logic, database uniqueness, and scheduling conflict rules.")
    add_bullet("5. Global Exception Handler (@RestControllerAdvice): ", "Intercepts all errors centrally and formats them into clean HTTP JSON responses.")
    add_bullet("6. Persistence Layer (Spring Data JPA & Hibernate): ", "Maps entities to MySQL database tables.")

    # 2. DTO Layer
    add_h1("2. DEEP-DIVE: DTO LAYER & INPUT BLUNDER CONSTRAINTS")
    add_body("DTOs (Data Transfer Objects) act as the security boundary of the API. When a user submits an HTTP request, the DTO inspects the payload before any business logic executes.")
    
    dto_headers = ["DTO File", "Field", "Validation Constraint", "Blunder Prevented", "Example Bad Input Caught"]
    dto_data = [
        ["RegisterRequest", "name", "@NotBlank, @Pattern(^[a-zA-Z\\s]{2,50}$)", "Numbers/symbols in full name", "'234', 'John@123'"],
        ["RegisterRequest", "phone", "@NotBlank, @Pattern(^[0-9]{10}$)", "Non-numeric or wrong length phone", "'sudhansh', '12345'"],
        ["RegisterRequest", "email", "@NotBlank, @Email", "Malformed email structure", "'not-an-email', 'user@'"],
        ["RegisterRequest", "password", "@NotBlank, @Size(min = 6)", "Insecure short password", "'123', 'abc'"],
        ["AppointmentRequest", "patientId", "@NotNull", "Missing patient reference ID", "null"],
        ["AppointmentRequest", "doctorId", "@NotNull", "Missing doctor reference ID", "null"],
        ["AppointmentRequest", "appointmentDate", "@NotBlank", "Missing / blank appointment date", "null, ''"],
        ["AppointmentRequest", "appointmentTime", "@NotBlank", "Missing / blank appointment time slot", "null, ''"]
    ]
    add_table(dto_headers, dto_data, [Inches(1.5), Inches(1.0), Inches(2.2), Inches(1.4), Inches(0.9)])

    # 3. Global Exception Handling
    add_h1("3. DEEP-DIVE: GLOBAL EXCEPTION HANDLING & ERROR RESPONSES")
    add_body("Without centralized exception handling, errors result in raw 500 server crashes with ugly Java stack traces. GlobalExceptionHandler (@RestControllerAdvice) intercepts all exceptions and wraps them into a standardized ErrorResponse object:")
    
    exc_headers = ["Exception Caught", "Trigger Source", "HTTP Status Returned", "Client Message Example"]
    exc_data = [
        ["MethodArgumentNotValidException", "DTO @Valid annotation failures", "400 Bad Request", "Phone must be a valid 10-digit number"],
        ["IllegalArgumentException", "Service layer validation / blunder checks", "400 Bad Request", "Cannot book an appointment for a past date: 2020-01-01"],
        ["AppointmentAlreadyExistsException", "Doctor double-booking slot conflict", "409 Conflict", "Doctor is already booked at this time."],
        ["RuntimeException", "Entity lookups & business errors", "400 Bad Request", "Patient not found / Invalid password"],
        ["Exception (Generic)", "Unexpected server failures", "500 Internal Server", "An unexpected error occurred"]
    ]
    add_table(exc_headers, exc_data, [Inches(2.2), Inches(1.8), Inches(1.1), Inches(1.9)])

    # 4. Service Layer Logic
    add_h1("4. DEEP-DIVE: SERVICE LAYER BUSINESS RULES")
    add_body("Rules that require database queries or date calculations are executed inside the @Service classes:")
    add_bullet("Duplicate Email Check (UserService): ", "Queries MySQL with userRepository.findByEmail(). If already exists, throws IllegalArgumentException.")
    add_bullet("Past Date Prevention (AppointmentService): ", "Compares LocalDate.parse(date) with LocalDate.now(). If past date, throws IllegalArgumentException.")
    add_bullet("Doctor Double-Booking Check (AppointmentService): ", "Queries appointmentRepository for matching doctor, date, and time. If exists, throws AppointmentAlreadyExistsException.")
    add_bullet("Patient Simultaneous Booking (AppointmentService): ", "Prevents the same patient from booking two doctors at the same hour.")
    add_bullet("Blank Prescription Check (AppointmentService): ", "Rejects whitespace-only clinical notes in updatePrescription().")

    # 5. Frontend Guards
    add_h1("5. DEEP-DIVE: FRONTEND BLUNDER GUARDS")
    add_bullet("Real-Time Phone Filter (Signup.js): ", "Strips non-digit characters on keystroke and limits input length to exactly 10 digits.")
    add_bullet("Pre-Submit Regex Check (Signup.js): ", "Validates letters-only name, email format, and password length, displaying instant alert banners.")
    add_bullet("Calendar Past Date Lock (Dashboard.js): ", "Sets min={today} on HTML5 date picker so past dates are disabled in the UI.")
    add_bullet("Prescription Modal Guard (Dashboard.js): ", "Prevents doctors from submitting empty prescription notes.")

    # 6. Request Lifecycle Walkthrough
    add_h1("6. END-TO-END REQUEST LIFECYCLE WALKTHROUGHS")
    add_h2("Scenario A: User Inputs Phone as 'sudhansh'")
    add_body("1. React form intercepts or sends JSON: { name: 'Rahul', phone: 'sudhansh', ... }.")
    add_body("2. UserController receives request; @Valid triggers DTO inspection.")
    add_body("3. RegisterRequest rejects phone: fails @Pattern(^[0-9]{10}$).")
    add_body("4. Spring throws MethodArgumentNotValidException.")
    add_body("5. GlobalExceptionHandler catches exception -> returns HTTP 400 with message 'Phone must be a valid 10-digit number'.")
    add_body("6. Database remains 100% clean.")

    add_h2("Scenario B: Booking for Past Date ('2020-01-01')")
    add_body("1. Request reaches AppointmentService.createAppointment().")
    add_body("2. LocalDate.parse('2020-01-01').isBefore(LocalDate.now()) evaluates to true.")
    add_body("3. Service throws IllegalArgumentException('Cannot book an appointment for a past date').")
    add_body("4. GlobalExceptionHandler catches it -> returns HTTP 400 Bad Request.")

    # 7. Unit Testing Strategy
    add_h1("7. UNIT TESTING STRATEGY (JUnit 5 & Mockito)")
    add_body("The test suite consists of 40 isolated unit tests across 4 service test classes:")
    add_bullet("AppointmentServiceTest (13 Tests): ", "Valid bookings, missing date/time, past date rejections, double-bookings, rescheduling, prescriptions.")
    add_bullet("UserServiceTest (15 Tests): ", "Valid registrations, default roles, duplicate emails, invalid names ('234'), invalid phones ('sudhansh'), bad emails, short passwords, auth login.")
    add_bullet("DoctorServiceTest (8 Tests): ", "Doctor creation, availability updates, fee updates, shift timings, error lookups.")
    add_bullet("PatientServiceTest (4 Tests): ", "Patient profile creation, listings, medical history updates, error lookups.")
    add_body("Result: 40 / 40 PASS (100% SUCCESS, Total Execution Time: 2.33 seconds, BUILD SUCCESS).", bold_prefix="Execution Result: ")

    try:
        doc.save(output_path)
        print(f"Architecture guide saved to: {output_path}")
    except PermissionError:
        print(f"Warning: {output_path} is currently open. Skipping.")

if __name__ == "__main__":
    build_architecture_doc(r"c:\Users\sudha\Desktop\healthcare\System_Architecture_and_Code_Flow_Guide.docx")
