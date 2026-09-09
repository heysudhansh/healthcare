import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_table_section(doc, heading_text, headers, data, col_widths):
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(14)
    p_head.paragraph_format.space_after = Pt(4)
    r = p_head.add_run(heading_text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header Row
    for c_idx, cell in enumerate(table.rows[0].cells):
        cell.width = col_widths[c_idx]
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(headers[c_idx])
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=45, bottom=45, left=80, right=80)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            val = str(row_data[c_idx])
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if c_idx == len(row_data) - 1:
                r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)

def build_simple_discrete_report(output_path):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("UNIT TESTING REPORT - HOSPITAL MANAGEMENT SYSTEM")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    r = p_sub.add_run("Backend Service Layer Test Execution and Input Validation Results")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0, 0, 0)

    # Overall Summary Table
    summary_headers = ["Total Tests", "Passed", "Failed", "Skipped", "Execution Time", "Build Status"]
    summary_data = [["40", "40", "0", "0", "2.33s", "BUILD SUCCESS"]]
    summary_widths = [Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.3), Inches(1.3)]
    add_table_section(doc, "OVERALL TEST SUMMARY", summary_headers, summary_data, summary_widths)

    col_widths_standard = [Inches(0.4), Inches(2.2), Inches(2.2), Inches(2.2), Inches(0.5)]
    std_headers = ["#", "Test Function Name", "Test Scenario / Input", "Expected Result", "Status"]

    # SECTION 1: APPOINTMENT SERVICE (13 Tests)
    appt_data = [
        ["1", "testCreateAppointment_Success", "Valid booking (Doctor ID, Patient ID, Future Date, Time)", "Appointment saved with status 'Booked' and consultation fee", "PASS"],
        ["2", "testCreateAppointment_PatientNotFound_ThrowsException", "Booking with non-existent Patient ID: 999", "Fails: Throws error 'Patient not found'", "PASS"],
        ["3", "testCreateAppointment_DoctorNotFound_ThrowsException", "Booking with non-existent Doctor ID: 999", "Fails: Throws error 'Doctor not found'", "PASS"],
        ["4", "testCreateAppointment_MissingDate_ThrowsException", "Booking with missing date (null)", "Fails: Throws error 'Appointment date is required.'", "PASS"],
        ["5", "testCreateAppointment_PastDate_ThrowsException", "Booking with past date (e.g. '2020-01-01')", "Fails: Throws error 'Cannot book an appointment for a past date'", "PASS"],
        ["6", "testCreateAppointment_MissingTime_ThrowsException", "Booking with missing time slot ('')", "Fails: Throws error 'Appointment time is required.'", "PASS"],
        ["7", "testCreateAppointment_DoctorDoubleBooking_ThrowsException", "Doctor already booked for same date and time slot", "Fails: Throws error 'Doctor is already booked at this time.'", "PASS"],
        ["8", "testCreateAppointment_PatientDoubleBooking_ThrowsException", "Patient already booked with another doctor at same time", "Fails: Throws error 'Patient already has an appointment at this time.'", "PASS"],
        ["9", "testCancelAppointment_Success", "Cancel existing appointment by ID: 501", "Appointment status updated to 'CANCELLED'", "PASS"],
        ["10", "testRescheduleAppointment_Success", "Reschedule to valid new future date and time slot", "Appointment updated with new date/time, status remains 'Booked'", "PASS"],
        ["11", "testRescheduleAppointment_PastDate_ThrowsException", "Reschedule to past date (e.g. '2020-01-01')", "Fails: Throws error 'Cannot reschedule appointment to a past date'", "PASS"],
        ["12", "testUpdatePrescription_Success", "Doctor writes and updates prescription notes", "Prescription saved successfully on appointment", "PASS"],
        ["13", "testUpdatePrescription_EmptyPrescription_ThrowsException", "Doctor submits blank / empty prescription (e.g. '   ')", "Fails: Throws error 'Prescription note cannot be empty.'", "PASS"],
    ]
    add_table_section(doc, "SECTION 1: APPOINTMENT SERVICE TESTS (AppointmentServiceTest - 13 Tests)", std_headers, appt_data, col_widths_standard)

    # SECTION 2: USER SERVICE (15 Tests)
    user_data = [
        ["1", "testRegisterPatient_Success", "Register new patient with valid details", "User saved with PATIENT role and patient record initialized", "PASS"],
        ["2", "testRegisterDoctor_WithShiftAndHours", "Register new doctor with shift timing and working hours", "User saved with DOCTOR role and doctor record initialized", "PASS"],
        ["3", "testRegisterDoctor_AppliesDefaultsWhenFieldsBlank", "Register doctor with blank specialization and qualification", "Applies default fallbacks ('General Physician', 'MBBS, MD')", "PASS"],
        ["4", "testRegisterUser_DefaultRolePatientWhenNull", "Register user with role as null", "Defaults user role to PATIENT", "PASS"],
        ["5", "testRegisterUser_InvalidName_ThrowsException", "Name containing numbers (e.g. '234')", "Fails: Throws error 'Invalid name. Name must contain only letters'", "PASS"],
        ["6", "testRegisterUser_InvalidPhone_ThrowsException", "Phone containing text (e.g. 'sudhansh')", "Fails: Throws error 'Invalid phone number. Must be a 10-digit number'", "PASS"],
        ["7", "testRegisterUser_InvalidEmail_ThrowsException", "Malformed email format (e.g. 'not-an-email')", "Fails: Throws error 'Invalid email format'", "PASS"],
        ["8", "testRegisterUser_ShortPassword_ThrowsException", "Password shorter than 6 characters (e.g. '123')", "Fails: Throws error 'Password must be at least 6 characters long'", "PASS"],
        ["9", "testRegisterUser_DuplicateEmail_ThrowsException", "Registering with an already registered email in DB", "Fails: Throws error 'Email is already registered'", "PASS"],
        ["10", "testSaveUser_PatientRole_CreatesPatientIfNotExists", "Save user with PATIENT role (creates profile if missing)", "Creates linked patient profile in database", "PASS"],
        ["11", "testSaveUser_DoctorRole_CreatesDoctorIfNotExists", "Save user with DOCTOR role (creates profile if missing)", "Creates linked doctor profile with default fee/hours", "PASS"],
        ["12", "testLogin_Success", "Login with valid registered email and password", "Returns matching authenticated user details", "PASS"],
        ["13", "testLogin_InvalidPassword_ThrowsException", "Login with incorrect password", "Fails: Throws error 'Invalid password'", "PASS"],
        ["14", "testLogin_UserNotFound_ThrowsException", "Login with unregistered email", "Fails: Throws error 'User not found'", "PASS"],
        ["15", "testGetAllUsers_ReturnsUserList", "Fetch all users list for admin", "Returns list of all registered users", "PASS"],
    ]
    add_table_section(doc, "SECTION 2: USER SERVICE & VALIDATION TESTS (UserServiceTest - 15 Tests)", std_headers, user_data, col_widths_standard)

    # SECTION 3: DOCTOR SERVICE (8 Tests)
    doc_data = [
        ["1", "testSaveDoctor_Success", "Save and create doctor entity", "Doctor profile saved with specialization and consultation fee", "PASS"],
        ["2", "testGetAllDoctors_ReturnsList", "Get list of all registered doctors", "Returns list containing all doctor profiles", "PASS"],
        ["3", "testUpdateAvailability_Success", "Update doctor availability status to 'Busy'", "Availability updated to 'Busy'", "PASS"],
        ["4", "testUpdateAvailability_DoctorNotFound_ThrowsException", "Update availability for non-existent doctor ID: 999", "Fails: Throws error 'Doctor not found'", "PASS"],
        ["5", "testUpdateConsultationFee_Success", "Update consultation fee to ₹2000", "Consultation fee updated to ₹2000", "PASS"],
        ["6", "testUpdateConsultationFee_DoctorNotFound_ThrowsException", "Update consultation fee for non-existent doctor ID: 999", "Fails: Throws error 'Doctor not found'", "PASS"],
        ["7", "testUpdateShiftAndWorkingHours_Success", "Update shift to 'Night' and working hours to 6", "Shift and daily working hours updated", "PASS"],
        ["8", "testUpdateShift_DoctorNotFound_ThrowsException", "Update shift for non-existent doctor ID: 999", "Fails: Throws error 'Doctor not found'", "PASS"],
    ]
    add_table_section(doc, "SECTION 3: DOCTOR SERVICE TESTS (DoctorServiceTest - 8 Tests)", std_headers, doc_data, col_widths_standard)

    # SECTION 4: PATIENT SERVICE (4 Tests)
    pat_data = [
        ["1", "testSavePatient_Success", "Save and create patient entity", "Patient profile saved with gender and blood group", "PASS"],
        ["2", "testGetAllPatients_ReturnsList", "Get list of all registered patients", "Returns list containing all patient records", "PASS"],
        ["3", "testUpdateMedicalHistory_Success", "Update patient medical history", "Medical history updated successfully", "PASS"],
        ["4", "testUpdateMedicalHistory_PatientNotFound_ThrowsException", "Update medical history for non-existent patient ID: 999", "Fails: Throws error 'Patient not found'", "PASS"],
    ]
    add_table_section(doc, "SECTION 4: PATIENT SERVICE TESTS (PatientServiceTest - 4 Tests)", std_headers, pat_data, col_widths_standard)

    try:
        doc.save(output_path)
        print(f"Report saved to: {output_path}")
    except PermissionError:
        print(f"Warning: {output_path} is currently open in another program. Skipping.")

if __name__ == "__main__":
    for path in [
        r"c:\Users\sudha\Desktop\healthcare\Unit_Test_Report.docx",
        r"c:\Users\sudha\Desktop\healthcare\Hospital_Unit_Test_Report.docx",
        r"c:\Users\sudha\Desktop\healthcare\Unit_Test_Report_v2.docx"
    ]:
        build_simple_discrete_report(path)
